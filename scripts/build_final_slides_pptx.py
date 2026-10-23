#!/usr/bin/env python3
"""
Compiles docs/final/slides.md into a professional Microsoft PowerPoint (.pptx) deck.
Outputs: docs/final/slides.pptx
"""

import sys
from pathlib import Path

MD_SLIDES = Path("docs/final/slides.md")
PPTX_FILE = Path("docs/final/slides.pptx")


def build_slides():
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor
        from pptx.enum.text import PP_ALIGN
    except ImportError:
        print("python-pptx not installed; skipping PPTX generation.")
        return

    prs = Presentation()
    # Set 16:9 widescreen dimensions (13.33 x 7.5 inches)
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]

    with open(MD_SLIDES, "r", encoding="utf-8") as f:
        content = f.read()

    slides_raw = content.split("---")

    for raw in slides_raw:
        lines = [l.strip() for l in raw.strip().split("\n") if l.strip()]
        if not lines:
            continue

        slide = prs.slides.add_slide(blank_layout)

        # Header Title
        title_text = "Slide"
        body_bullets = []

        for line in lines:
            if line.startswith("# ") or line.startswith("## "):
                title_text = line.lstrip("#").strip()
            else:
                body_bullets.append(line)

        # Title box
        txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.7), Inches(1.0))
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = RGBColor(26, 37, 47)

        # Body box
        bodyBox = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
        btf = bodyBox.text_frame
        btf.word_wrap = True

        for idx, bullet in enumerate(body_bullets):
            p_bullet = btf.add_paragraph() if idx > 0 else btf.paragraphs[0]
            clean_b = bullet.lstrip("-* ").strip()
            if bullet.startswith("- **") or bullet.startswith("**"):
                p_bullet.font.bold = True
                p_bullet.font.size = Pt(17)
                p_bullet.font.color.rgb = RGBColor(41, 128, 185)
            else:
                p_bullet.font.size = Pt(15)
                p_bullet.font.color.rgb = RGBColor(44, 62, 80)
            p_bullet.text = clean_b

    prs.save(PPTX_FILE)
    print(f"[+] Successfully compiled PowerPoint Deck: {PPTX_FILE}")


if __name__ == "__main__":
    build_slides()
