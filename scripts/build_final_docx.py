#!/usr/bin/env python3
"""
Compiles docs/final/Final_Report.md into a formatted Microsoft Word (.docx) document.
Outputs: docs/final/Final_Report.docx
"""

from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

MD_FILE = Path("docs/final/Final_Report.md")
DOCX_FILE = Path("docs/final/Final_Report.docx")


def build_docx():
    if not MD_FILE.exists():
        print(f"Error: {MD_FILE} does not exist.")
        return

    doc = docx.Document()

    # Set page margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    with open(MD_FILE, "r", encoding="utf-8") as f:
        lines = f.readlines()

    in_code_block = False
    code_text = []

    for line in lines:
        stripped = line.strip()

        # Handle Code blocks
        if stripped.startswith("```"):
            if in_code_block:
                in_code_block = False
                p = doc.add_paragraph("\n".join(code_text))
                p.style.font.name = "Consolas"
                p.style.font.size = Pt(9.5)
                code_text = []
            else:
                in_code_block = True
                code_text = []
            continue

        if in_code_block:
            code_text.append(line.rstrip("\n"))
            continue

        # Handle Headers
        if stripped.startswith("# "):
            h = doc.add_heading(stripped[2:], level=1)
            h.alignment = WD_ALIGN_PARAGRAPH.LEFT
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=2)
        elif stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=3)
        elif stripped.startswith("#### "):
            doc.add_heading(stripped[5:], level=4)
        elif stripped.startswith("- ") or stripped.startswith("* "):
            doc.add_paragraph(stripped[2:], style="List Bullet")
        elif stripped.startswith("1. ") or stripped.startswith("2. ") or stripped.startswith("3. ") or stripped.startswith("4. "):
            doc.add_paragraph(stripped[3:], style="List Number")
        elif stripped.startswith("|"):
            # Table row or separator
            if "---" in stripped:
                continue
            cells = [c.strip() for c in stripped.split("|")[1:-1]]
            p = doc.add_paragraph("  |  ".join(cells))
            p.style.font.name = "Consolas"
            p.style.font.size = Pt(9.0)
        elif stripped:
            doc.add_paragraph(stripped)

    doc.save(DOCX_FILE)
    print(f"[+] Successfully compiled Word Document: {DOCX_FILE}")


if __name__ == "__main__":
    build_docx()
